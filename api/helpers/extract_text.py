import io

import fitz  # PyMuPDF
import openpyxl
from docx import Document as DocxDocument

from models.hot_store import FileType


def extract_text(data: bytes, file_type: FileType) -> str:
    if file_type == FileType.PDF:
        doc = fitz.open(stream=data, filetype="pdf")
        return "\n".join(page.get_text() for page in doc)

    if file_type == FileType.DOCX:
        doc = DocxDocument(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs)

    if file_type == FileType.XLSX:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines: list[str] = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                lines.append("\t".join(str(c) if c is not None else "" for c in row))
        return "\n".join(lines)

    if file_type in (FileType.CSV, FileType.TXT):
        return data.decode("utf-8", errors="replace")

    return ""
