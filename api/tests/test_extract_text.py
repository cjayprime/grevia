"""Unit tests for helpers/extract_text.py — uses in-memory file bytes."""
import io
import csv

import pytest

from models.hot_store import FileType
from helpers.extract_text import extract_text


def _make_csv(rows: list[list]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerows(rows)
    return buf.getvalue().encode()


def _make_docx(paragraphs: list[str]) -> bytes:
    from docx import Document
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_xlsx(rows: list[list]) -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestExtractText:
    def test_txt_returns_content(self):
        data = b"Hello, world!"
        assert extract_text(data, FileType.TXT) == "Hello, world!"

    def test_csv_returns_rows(self):
        data = _make_csv([["scope", "tco2e"], ["1", "100"], ["2", "200"]])
        result = extract_text(data, FileType.CSV)
        assert "scope" in result
        assert "100" in result

    def test_docx_returns_paragraphs(self):
        data = _make_docx(["First paragraph", "Second paragraph"])
        result = extract_text(data, FileType.DOCX)
        assert "First paragraph" in result
        assert "Second paragraph" in result

    def test_xlsx_returns_cell_values(self):
        data = _make_xlsx([["Name", "Value"], ["Emissions", 500]])
        result = extract_text(data, FileType.XLSX)
        assert "Emissions" in result
        assert "500" in result

    def test_other_returns_empty_string(self):
        assert extract_text(b"\x00\x01\x02", FileType.OTHER) == ""

    def test_txt_utf8_decoding(self):
        data = "Ångström units — €500".encode("utf-8")
        result = extract_text(data, FileType.TXT)
        assert "€500" in result

    def test_empty_txt(self):
        assert extract_text(b"", FileType.TXT) == ""

    def test_docx_empty_document(self):
        data = _make_docx([])
        result = extract_text(data, FileType.DOCX)
        assert isinstance(result, str)
